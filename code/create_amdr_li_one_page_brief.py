"""Create a one-page, beginner-friendly AMDR-Li study brief."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs/briefs/amdr_li_beginner_study_brief.docx")
NAVY = "17365D"
BLUE = "2E74B5"
INK = "202124"
MUTED = "5F6368"
LIGHT_BLUE = "E8EEF5"
LIGHT_GOLD = "FFF5D8"
WHITE = "FFFFFF"
WIDTH_DXA = 9360


def set_font(run, size=10.0, color=INK, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def fixed_width_table(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = int(width)
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def paragraph(doc, text="", before=0, after=4, size=10.0, color=INK, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def label_paragraph(doc, label, text, before=0, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(label)
    set_font(run, size=9.6, color=NAVY, bold=True)
    run = p.add_run(text)
    set_font(run, size=9.6)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=11.5, color=BLUE, bold=True)
    return p


def add_callout(doc):
    table = doc.add_table(rows=1, cols=1)
    fixed_width_table(table, [WIDTH_DXA])
    cell = table.cell(0, 0)
    shade(cell, LIGHT_GOLD)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run("Bottom line. ")
    set_font(r, size=9.7, color=NAVY, bold=True)
    r = p.add_run(
        "AMDR-Li is the strongest model on RMSE, meaning it reduces the largest forecasting misses. "
        "It is not yet a proven real-time trading system or a universal winner on every metric."
    )
    set_font(r, size=9.7)
    paragraph(doc, "", after=0, size=1)


def add_scorecard(doc):
    heading(doc, "4. How good is it - and what still needs work")
    table = doc.add_table(rows=1, cols=4)
    fixed_width_table(table, [1960, 1850, 2750, 2800])
    headers = ["Target", "RMSE result", "Where it is weaker", "Trading / publication reading"]
    for c, text in zip(table.rows[0].cells, headers):
        shade(c, LIGHT_BLUE)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=8.2, color=NAVY, bold=True)
    rows = [
        ("Carbonate", "Best RMSE\n2,881", "MAE, MASE, MAPE, sMAPE and direction trail the best comparators.", "AMDR-Li Scheme 1: Sharpe 7.81, but VMD-ARIMA is stronger (11.52)."),
        ("Hydroxide", "Best RMSE\n0.274", "Direction is only tied-best (44.29%); Random Walk has lower MAE and MAPE.", "Sharpe 4.56; CEEMDAN-ARIMA is marginally higher (4.68)."),
        ("ALB equity", "Best RMSE\n2.328", "No clear point-metric loss within the tested set.", "Strong extension; 53.67% lower RMSE than Random Walk."),
        ("Ganfeng", "Best RMSE\n0.190", "VMD-ARIMA has better MAE, MASE, MAPE and sMAPE.", "Useful but narrow RMSE gain: 0.55% over VMD-ARIMA."),
    ]
    for values in rows:
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i in (0, 1):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_font(r, size=7.65, color=INK, bold=(i == 0))
    p = paragraph(doc, "Trading caution: the Sharpe ratios, returns and drawdowns are descriptive only. Because the study is retrospective, they cannot yet be treated as investable performance or evidence of abnormal return.", before=3, after=0, size=8.6, color=MUTED, italic=True)
    return table


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.60)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    sec.header_distance = Inches(0.30)
    sec.footer_distance = Inches(0.30)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(9.6)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    # Named one-page override: denser than the standard-business-brief preset,
    # while preserving its type hierarchy and restrained palette.
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("AMDR-Li research brief")
    set_font(r, size=8, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Lithium forecasting study | September 2026")
    set_font(r, size=8, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("RESEARCH NOTE")
    set_font(r, size=8.5, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("AMDR-Li: what the results actually mean")
    set_font(r, size=18, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("A plain-English one-page summary of the core-paper replication and the lithium extension")
    set_font(r, size=9.5, color=MUTED, italic=True)

    add_callout(doc)

    heading(doc, "1. Why an 80:20 holdout is not real-world implementation")
    paragraph(doc, "Imagine 100 historical trading days. A static 80:20 study builds the model using the first 80 days, then reports results on the final 20. That is useful as a laboratory comparison, but a live forecaster must make every decision using only information known before each next-day forecast.", after=3, size=9.55)
    label_paragraph(doc, "What both studies lack: ", "The core paper and our current AMDR-Li paper share this limitation. Full-sample decomposition, global preparation and observed-test ARIMA filtering use later observations. That makes the headline results retrospective evidence, not a simulation of decisions made day by day in real time.")

    heading(doc, "2. What happened when we replicated the fixed older method")
    paragraph(doc, "ARIMA is a statistical model that projects the next value from a component's recent history; it is useful for smoother, more linear patterns. LSTM is a neural-network sequence model; it can learn more complex nonlinear patterns but needs more data and tuning. The fixed older recipe always sends the first low-frequency components to ARIMA, all remaining components to LSTM, then adds the forecasts. On lithium, that rigid rule did not produce a universal win.", after=3, size=9.35)

    heading(doc, "3. What AMDR-Li changes")
    label_paragraph(doc, "The actual difference: ", "AMDR-Li keeps the same broad decomposition idea, but replaces the core paper's fixed routing rule. It tests VMD and CEEMDAN candidates, selects features separately for each component with LASSO, then lets development-period error choose ARIMA, LSTM, SVR or Random Forest for each component. It combines the strongest candidate forecasts with weights fixed before the final 20% is scored.")
    label_paragraph(doc, "Multimodal means: ", "it uses lagged prices plus linked lithium prices, producer equities, lithium-specific search interest and news sentiment, sector indices and market controls. Li simply identifies the lithium application.")

    add_scorecard(doc)
    label_paragraph(doc, "Publication judgement: ", "This can support a transparent retrospective methodology paper, but it is not yet sufficient for a live forecasting or trading claim. We did run a past-only walk-forward test on the earlier Glencore critical-minerals pipeline: its hybrid RMSE was 0.240 versus 0.112 for Random Walk, so the advantage did not survive. Current AMDR-Li lithium targets still need that equivalent causal test, plus benchmark-relative return and factor-adjusted risk evidence.", before=3, after=0)
    doc.core_properties.title = "AMDR-Li Beginner Study Brief"
    doc.core_properties.author = "Pratham Lahoti"
    doc.core_properties.subject = "Plain-English summary of retrospective lithium forecasting evidence"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
