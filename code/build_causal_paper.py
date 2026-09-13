"""Build a publication-style DOCX from freshly generated causal paper assets.

Run this with the bundled workspace Python runtime after
``generate_paper_assets.py``.  It uses the ``narrative_proposal`` document
preset with an editorial-cover opening and one named landscape override for
the two wide trading tables.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "555555"
TABLE_FILL = "F4F6F9"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
PROPOSED = "VMD-LASSO-ARX/LSTM"


def _set_font(run, *, size: float, color: str = "000000", bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = TABLE_INDENT_DXA) -> None:
    """Apply fixed DXA widths to table, grid, and all cells."""

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths_dxa):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar"); fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar"); fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1); run._r.append(instr); run._r.append(fld_char2)


def _configure_section(section, *, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11), Inches(8.5)
        section.top_margin = section.bottom_margin = Inches(0.72)
        section.left_margin = section.right_margin = Inches(0.75)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def _setup_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for style_name, size, color, before, after in (("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK_BLUE, 8, 4)):
        style = document.styles[style_name]
        style.font.name = "Calibri"; style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)


def _header_footer(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("Critical-Mineral Equity Forecasting | Causal Empirical Study")
    _set_font(run, size=8.5, color=MUTED)
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Causal one-step evaluation | Page ")
    _set_font(run, size=8.5, color=MUTED)
    _add_page_field(paragraph)


def _add_paragraph(document: Document, text: str, *, bold_prefix: str | None = None, italic: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.333
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix); _set_font(prefix, size=11, bold=True)
        remainder = paragraph.add_run(text[len(bold_prefix):]); _set_font(remainder, size=11, italic=italic)
    else:
        run = paragraph.add_run(text); _set_font(run, size=11, italic=italic)


def _add_caption(document: Document, kind: str, number: int, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4); paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    label = paragraph.add_run(f"{kind} {number}. "); _set_font(label, size=9, bold=True, color=DARK_BLUE)
    body = paragraph.add_run(caption); _set_font(body, size=9, italic=True)


def _column_widths(frame: pd.DataFrame, total_dxa: int) -> list[int]:
    measures = []
    for column in frame.columns:
        visible = [str(column)] + [str(value) for value in frame[column].head(20)]
        measures.append(min(34, max(7, max(len(value) for value in visible))))
    total = sum(measures)
    widths = [max(420, int(total_dxa * measure / total)) for measure in measures]
    widths[-1] += total_dxa - sum(widths)
    return widths


def _add_dataframe(document: Document, frame: pd.DataFrame, *, landscape: bool = False, font_size: float | None = None) -> None:
    text = frame.copy().fillna("")
    for column in text.columns:
        text[column] = text[column].map(lambda value: f"{value:.4f}" if isinstance(value, float) else str(value))
    table = document.add_table(rows=1, cols=len(text.columns))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    header = table.rows[0]
    _set_repeat_table_header(header)
    for index, column in enumerate(text.columns):
        cell = header.cells[index]; cell.text = ""
        paragraph = cell.paragraphs[0]; paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(column)); _set_font(run, size=font_size or (6.8 if landscape else 7.5), color=NAVY, bold=True)
        _shade(cell, TABLE_FILL); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in text.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]; cell.text = ""; cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]; paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value); _set_font(run, size=font_size or (6.8 if landscape else 7.5))
            paragraph.paragraph_format.space_after = Pt(0); paragraph.paragraph_format.line_spacing = 1.0
    content_width = 13680 if landscape else TABLE_WIDTH_DXA
    _set_table_geometry(table, _column_widths(text, content_width), indent_dxa=120 if not landscape else 90)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def _add_figure(document: Document, path: Path, number: int, caption: str) -> None:
    _add_caption(document, "Figure", number, caption)
    paragraph = document.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(); run.add_picture(str(path), width=Inches(6.25))


def _narrative_values(tables: dict[int, pd.DataFrame]) -> tuple[str, str, str]:
    single = tables[7].set_index("Model")
    decomp = tables[8].set_index("Model")
    random_walk = float(single.loc["Random Walk", "RMSE"]) if "Random Walk" in single.index else float("nan")
    arima = float(single.loc["ARIMA", "RMSE"]) if "ARIMA" in single.index else float("nan")
    proposed = float(decomp.loc[PROPOSED, "RMSE"]) if PROPOSED in decomp.index else float("nan")
    return f"{random_walk:.4f}", f"{arima:.4f}", f"{proposed:.4f}"


def build(asset_dir: Path, output_docx: Path) -> Path:
    asset_dir, output_docx = Path(asset_dir).resolve(), Path(output_docx).resolve()
    manifest = json.loads((asset_dir / "asset_manifest.json").read_text(encoding="utf-8"))
    tables = {number: pd.read_csv(asset_dir / "tables" / f"table{number}.csv") for number in range(1, 13)}
    document = Document(); _configure_section(document.sections[0]); _setup_styles(document); _header_footer(document.sections[0])
    # Editorial-cover opening under the narrative_proposal preset.
    document.add_paragraph().paragraph_format.space_after = Pt(42)
    kicker = document.add_paragraph(); kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = kicker.add_run("RESEARCH MANUSCRIPT"); _set_font(r, size=11, color=DARK_BLUE, bold=True)
    title = document.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_before = Pt(10); title.paragraph_format.space_after = Pt(10); r = title.add_run("Causal Multimodal Forecasting and Trading Evaluation for Critical-Mineral Equities"); _set_font(r, size=25, color=NAVY, bold=True)
    subtitle = document.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER; subtitle.paragraph_format.space_after = Pt(30); r = subtitle.add_run("A transparent one-step walk-forward assessment across eight global companies"); _set_font(r, size=13, color=MUTED, italic=True)
    meta = document.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = meta.add_run(f"Canonical run: {manifest.get('canonical_run_id')}\nPrepared: {manifest.get('created_utc', '')[:10]}"); _set_font(r, size=10, color=MUTED)
    disclosure = document.add_paragraph(); disclosure.alignment = WD_ALIGN_PARAGRAPH.CENTER; disclosure.paragraph_format.space_before = Pt(38); r = disclosure.add_run("Research-integrity note: all principal results use a causal design. The noncausal released-paper shadow is excluded from the empirical findings."); _set_font(r, size=9.5, color="9B1C1C", italic=True)
    document.add_page_break()
    document.add_heading("Abstract", level=1)
    rw, arima, proposed = _narrative_values(tables)
    _add_paragraph(document, "This paper evaluates one-day-ahead forecasts and economically motivated trading signals for eight global critical-mineral equities. The design uses a chronological final holdout, a calibration segment that precedes it, and scheduled model re-estimation that consumes only information available at each forecast origin. Candidate models include statistical, regularized, tree-based, recurrent, and decomposition-based approaches. All prediction intervals are sequential conformal intervals and all trading results deduct specified transaction costs. Across companies, the mean final-test RMSE is reported transparently for the random walk, ARIMA, and the proposed VMD-LASSO-ARX/LSTM architecture rather than inferred from a static split. The manuscript therefore emphasizes empirical comparison and limitations over a claim of universal model superiority.")
    _add_paragraph(document, "Keywords: critical minerals; equity forecasting; causal evaluation; variational mode decomposition; LASSO; conformal prediction; transaction costs.", bold_prefix="Keywords:")
    document.add_heading("1. Introduction", level=1)
    _add_paragraph(document, "Critical-mineral producers are exposed to commodity cycles, market stress, currency conditions, sector benchmarks, public attention, and information tone. These channels motivate multimodal predictor sets, but they also create a substantial risk of look-ahead bias when features, decompositions, or model choices are fitted on a final test period. This study treats causal timing as a design constraint rather than a post-hoc diagnostic.")
    _add_paragraph(document, "The empirical contribution is an auditable eight-company one-step evaluation that records the origin date, model-fit date, forecast, interval, and trading signal for every final-test observation. It provides a direct comparison with both simple benchmarks and complex decomposition models under the same timing convention.")
    document.add_heading("2. Literature and research positioning", level=1)
    _add_caption(document, "Table", 1, "Representative research and methodological references informing the study."); _add_dataframe(document, tables[1])
    _add_figure(document, asset_dir / "figures" / "fig1_publication_trend.png", 1, "Publication-record trend for a narrowly defined Crossref title query. This is a descriptive search snapshot, not a systematic review.")
    _add_paragraph(document, "The methodological reference integrates VMD, LASSO-selected multimodal inputs, ARIMA/LSTM forecasting, and trading analysis. This study retains the broad modelling motivation but changes the target to listed critical-mineral companies and, crucially, uses train-only decomposition and one-step walk-forward evaluation.")
    document.add_heading("3. Methodology", level=1)
    _add_figure(document, asset_dir / "figures" / "fig2_lstm_structure.png", 2, "Conceptual structure of the LSTM cell used as one candidate forecaster.")
    _add_figure(document, asset_dir / "figures" / "fig3_forecasting_framework.png", 3, "Causal one-step framework. Every operation at a forecast origin has access only to information available before the target close.")
    _add_caption(document, "Table", 2, "Actual parameter settings used in the causal empirical run."); _add_dataframe(document, tables[2])
    _add_caption(document, "Table", 3, "Predictors used in the causal multimodal feature frame."); _add_dataframe(document, tables[3])
    _add_paragraph(document, "For each target date t, the feature frame uses lagged target information and exogenous values available by t-1. Parameters and full decompositions are re-estimated only at the configured refit cadence. Between those dates, the decomposition state is advanced causally with newly observed history. The final 20% is reserved as a chronological test period, while the preceding calibration segment supports residual-based intervals and threshold selection only.")
    document.add_heading("4. Data and exploratory evidence", level=1)
    _add_caption(document, "Table", 4, "Descriptive statistics for the eight company closing-price series."); _add_dataframe(document, tables[4])
    _add_figure(document, asset_dir / "figures" / "fig4_price_series.png", 4, "Normalized closing-price paths for the eight critical-mineral equities.")
    _add_figure(document, asset_dir / "figures" / "fig5_selected_predictor_wordcloud.png", 5, "Word cloud of predictors repeatedly selected by the causal ARX procedure; it is not a raw-news-headline word cloud.")
    _add_figure(document, asset_dir / "figures" / "fig6_sentiment.png", 6, "SF Fed news-sentiment index during the sample and its empirical distribution.")
    document.add_heading("5. Decomposition and causal feature selection", level=1)
    _add_figure(document, asset_dir / "figures" / "fig7_vmd_components.png", 7, "Illustrative VMD components fitted only to Glencore's pre-test training segment; the display is not a full-sample decomposition.")
    _add_figure(document, asset_dir / "figures" / "fig8_approximate_entropy.png", 8, "Approximate-entropy values for the Figure 7 training-only VMD components.")
    _add_caption(document, "Table", 5, "Descriptive characteristics of training-only Glencore VMD components."); _add_dataframe(document, tables[5])
    _add_caption(document, "Table", 6, "Most frequently selected ARX predictors across scheduled causal refits."); _add_dataframe(document, tables[6])
    _add_paragraph(document, "The decomposition display is descriptive. The actual forecasting pipeline repeats train-only decomposition at scheduled origins, so no component shown or fitted in the final test is derived from later test observations. Feature-selection rates summarize the causal selections recorded across model refits.")
    document.add_heading("6. Forecasting results", level=1)
    _add_figure(document, asset_dir / "figures" / "fig9_single_model_forecasts.png", 9, "Illustrative final-test forecasts for Glencore using single-model candidates.")
    _add_figure(document, asset_dir / "figures" / "fig10_error_barplots.png", 10, "Cross-company mean final-test RMSE by model, normalized to the random-walk mean RMSE.")
    _add_caption(document, "Table", 7, "Final-test forecasting errors for single-model candidates, averaged across companies."); _add_dataframe(document, tables[7])
    _add_caption(document, "Table", 8, "Final-test forecasting errors for decomposition-based candidates, averaged across companies."); _add_dataframe(document, tables[8])
    _add_caption(document, "Table", 9, "Diebold-Mariano comparisons of the proposed model with each benchmark across companies."); _add_dataframe(document, tables[9])
    _add_paragraph(document, f"The cross-company mean RMSEs are {rw} for the random walk, {arima} for ARIMA, and {proposed} for the proposed decomposition architecture. These averages are descriptive across heterogeneous price scales; the within-company rankings and the company-level DM results should therefore be read alongside them. A negative DM statistic favours the proposed model because the loss difference is defined as proposed squared loss minus benchmark squared loss.")
    document.add_heading("7. Interval forecasts and trading evaluation", level=1)
    _add_figure(document, asset_dir / "figures" / "fig11_interval_forecasts.png", 11, "Illustrative conformal interval forecasts for the proposed causal model on Glencore's final test set.")
    _add_figure(document, asset_dir / "figures" / "fig12_trading_strategy.png", 12, "Timing of the causal one-step trading evaluation.")
    _add_figure(document, asset_dir / "figures" / "fig13_trading_evaluation.png", 13, "Cross-company mean Sharpe ratios for directional strategies at 10 bp transaction costs.")
    _add_caption(document, "Table", 10, "Final-test conformal interval quality, averaged across companies."); _add_dataframe(document, tables[10])
    _add_paragraph(document, "Trading results are evaluated at zero, 10, 25, and 50 basis points. Signals are formed at the stored forecast origin and executed against the next realized return; the uncertainty-adjusted version may abstain when the conformal interval is too wide. Thresholds are selected on calibration observations before the final test.")
    landscape = document.add_section(WD_SECTION.NEW_PAGE); _configure_section(landscape, landscape=True); _header_footer(landscape)
    document.add_heading("7.1 Trading-performance schedules", level=1)
    _add_caption(document, "Table", 11, "Trading performance of decomposition-model forecasts under calibrated strategies, averaged across companies."); _add_dataframe(document, tables[11], landscape=True, font_size=6.2)
    _add_caption(document, "Table", 12, "Trading performance of single-model forecasts under calibrated strategies at 10 bp, averaged across companies. All cost schedules are provided in the appendix ledger."); _add_dataframe(document, tables[12], landscape=True, font_size=6.2)
    portrait = document.add_section(WD_SECTION.NEW_PAGE); _configure_section(portrait); _header_footer(portrait)
    document.add_heading("8. Discussion and limitations", level=1)
    _add_paragraph(document, "The empirical record does not support a blanket conclusion that a complex decomposed neural architecture dominates simple forecasts. Performance differs by company and metric, while transaction costs can materially change the attractiveness of a directional strategy. This is why the report retains random walk and ARIMA benchmarks, reports intervals and DM comparisons, and discloses factor-adjusted trading estimates in the companion ledger.")
    _add_paragraph(document, "The analysis uses closing prices and daily proxy variables. It does not model tradability, market impact, borrow constraints, asynchronous disclosures, or company-specific corporate actions beyond the supplied price histories. Factor coverage is also frequency-dependent for emerging-market assets. These are limitations, not evidence of zero risk.")
    document.add_heading("9. Conclusion", level=1)
    _add_paragraph(document, "This paper provides a fully traceable causal framework for evaluating multimodal critical-mineral equity forecasts. Its value is the integrity of the timing, model comparison, interval evaluation, and cost-aware trading analysis. The attached ledgers preserve the complete company/model metrics needed to reproduce or challenge each summary result.")
    document.add_heading("References", level=1)
    references = [
        "Diebold, F.X. and Mariano, R.S. (1995). Comparing predictive accuracy. Journal of Business & Economic Statistics, 13(3), 253-263.",
        "Dragomiretskiy, K. and Zosso, D. (2014). Variational mode decomposition. IEEE Transactions on Signal Processing, 62(3), 531-544.",
        "Hochreiter, S. and Schmidhuber, J. (1997). Long short-term memory. Neural Computation, 9(8), 1735-1780.",
        "Liu, S., Li, M., Yang, K., Wei, Y. and Wang, S. (2025). From forecasting to trading: A multimodal-data-driven approach to reversing carbon market losses. Energy Economics, 144, 108350. https://doi.org/10.1016/j.eneco.2025.108350.",
        "Tibshirani, R. (1996). Regression shrinkage and selection via the lasso. Journal of the Royal Statistical Society: Series B, 58(1), 267-288.",
        "Data sources: project-provided price and MVIS workbooks; Yahoo Finance macro closes; Google Trends; Federal Reserve Bank of San Francisco news sentiment; and regional Fama-French factor files. Exact URLs, retrieval times, and SHA-256 checksums are in the canonical run registry.",
    ]
    for item in references:
        paragraph = document.add_paragraph(style="Normal"); paragraph.paragraph_format.left_indent = Inches(.25); paragraph.paragraph_format.first_line_indent = Inches(-.25); paragraph.paragraph_format.space_after = Pt(4); run = paragraph.add_run(item); _set_font(run, size=9.5)
    document.add_heading("Appendix: complete empirical ledgers and reproducibility", level=1)
    _add_paragraph(document, "The companion appendix_ledgers directory contains the full, unaggregated final-test forecast metrics, interval metrics, DM comparisons, trading metrics for every company/model/strategy/cost combination, and feature-selection records. The asset manifest records the causal run identifier, input checksums, figure/table captions, and the versioned Crossref search snapshot used in Figure 1.")
    _add_paragraph(document, "The released-paper shadow output is intentionally absent from these ledgers and from the reported findings because it is a noncausal forensic reconstruction rather than a valid forecast evaluation.")
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    return output_docx


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--assets", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(build(args.assets, args.output))


if __name__ == "__main__":
    main()
